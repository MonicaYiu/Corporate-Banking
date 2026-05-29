import os, json, uuid, base64, re
from pathlib import Path
from openai import OpenAI
import pdfplumber
import io

from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, HTMLResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

app = FastAPI(title="CIMB Document Intelligence API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

BASE_DIR   = Path(__file__).parent.parent
UPLOAD_DIR = BASE_DIR / "uploads"
RULES_DIR  = BASE_DIR / "rules"
FRONTEND   = BASE_DIR / "frontend"
UPLOAD_DIR.mkdir(exist_ok=True)
RULES_DIR.mkdir(exist_ok=True)

client = OpenAI(api_key=os.environ.get("OPENAI_API_KEY", ""))
MODEL  = "gpt-4o"

# ─── Checklists ───────────────────────────────────────────────────────────────

CHECKLISTS = {
    "FlexiPay": [
        {"id":"c1","name":"公司注册证明（ACRA BizFile）","severity":"blocking","bucket":"主体"},
        {"id":"c2","name":"联系人基本资料","severity":"blocking","bucket":"主体"},
        {"id":"c3","name":"近12个月银行流水PDF（非CIMB客户，最多4家银行）","severity":"blocking","bucket":"财务"},
        {"id":"c4","name":"CIMB Business Account开户文件（现有客户）","severity":"blocking","bucket":"主体"},
        {"id":"c5","name":"Keyman身份证明（NRIC）","severity":"blocking","bucket":"主体"},
    ],
    "BizGrow": [
        {"id":"b1","name":"SME申请表（常规版/Islamic版）","severity":"blocking","bucket":"主体"},
        {"id":"b2","name":"公司注册证明（ACRA BizFile）","severity":"blocking","bucket":"主体"},
        {"id":"b3","name":"股权结构证明（≥51%本地股权）","severity":"blocking","bucket":"主体"},
        {"id":"b4","name":"前五大供应商/客户清单","severity":"advisory","bucket":"业务"},
        {"id":"b5","name":"申请金额、期限、用途说明","severity":"blocking","bucket":"主体"},
        {"id":"b6","name":"其他银行借款明细","severity":"blocking","bucket":"财务"},
        {"id":"b7","name":"Keyman身份证明（NRIC）","severity":"blocking","bucket":"主体"},
        {"id":"b8","name":"担保人资料","severity":"advisory","bucket":"担保"},
        {"id":"b9","name":"声明与授权书","severity":"blocking","bucket":"主体"},
    ],
    "BizAssist": [
        {"id":"a1","name":"SME BizGrow/BizAssist申请表（常规版/Islamic版）","severity":"blocking","bucket":"主体"},
        {"id":"a2","name":"公司注册证明（ACRA BizFile）","severity":"blocking","bucket":"主体"},
        {"id":"a3","name":"本地股权证明（≥30%）","severity":"blocking","bucket":"主体"},
        {"id":"a4","name":"Borrower group架构图/关联公司披露","severity":"blocking","bucket":"主体"},
        {"id":"a5","name":"集团财务报表（收入≤S$100m或雇员≤200）","severity":"blocking","bucket":"财务"},
        {"id":"a6","name":"申请金额、期限、用途说明","severity":"blocking","bucket":"主体"},
        {"id":"a7","name":"其他银行借款及担保明细","severity":"blocking","bucket":"财务"},
        {"id":"a8","name":"Keyman/担保人资料","severity":"blocking","bucket":"担保"},
        {"id":"a9","name":"声明与授权书","severity":"blocking","bucket":"主体"},
    ],
    "BizProp": [
        {"id":"p1","name":"SME物业申请表（常规版/Islamic版）","severity":"blocking","bucket":"主体"},
        {"id":"p2","name":"公司注册证明（ACRA BizFile）","severity":"blocking","bucket":"主体"},
        {"id":"p3","name":"物业信息表（地址、类型、面积、价格、tenure、用途）","severity":"blocking","bucket":"物业"},
        {"id":"p4","name":"物业估价报告","severity":"blocking","bucket":"物业"},
        {"id":"p5","name":"物业产权证书（Title Deed）","severity":"blocking","bucket":"物业"},
        {"id":"p6","name":"公司财务报表/现金流证明","severity":"blocking","bucket":"财务"},
        {"id":"p7","name":"其他银行借款明细","severity":"blocking","bucket":"财务"},
        {"id":"p8","name":"担保人资料","severity":"advisory","bucket":"担保"},
        {"id":"p9","name":"董事会/合伙人授权决议（Board Resolution）","severity":"blocking","bucket":"主体"},
    ],
    "BizPropPlus": [
        {"id":"pp1","name":"SME物业申请表（常规版/Islamic版）","severity":"blocking","bucket":"主体"},
        {"id":"pp2","name":"公司注册证明（ACRA BizFile）","severity":"blocking","bucket":"主体"},
        {"id":"pp3","name":"物业信息表（地址、类型、面积、价格、tenure、用途）","severity":"blocking","bucket":"物业"},
        {"id":"pp4","name":"物业估价报告","severity":"blocking","bucket":"物业"},
        {"id":"pp5","name":"物业产权证书（Title Deed）","severity":"blocking","bucket":"物业"},
        {"id":"pp6","name":"公司财务报表/现金流证明","severity":"blocking","bucket":"财务"},
        {"id":"pp7","name":"其他银行借款明细","severity":"blocking","bucket":"财务"},
        {"id":"pp8","name":"担保人资料","severity":"advisory","bucket":"担保"},
        {"id":"pp9","name":"董事会/合伙人授权决议（Board Resolution）","severity":"blocking","bucket":"主体"},
        {"id":"pp10","name":"超额抵押安排说明文件（>100% MOF部分）","severity":"blocking","bucket":"担保"},
    ],
    "BizAssure": [
        {"id":"as1","name":"SME申请表","severity":"blocking","bucket":"主体"},
        {"id":"as2","name":"公司注册证明（ACRA BizFile）","severity":"blocking","bucket":"主体"},
        {"id":"as3","name":"人寿保险保单原件","severity":"blocking","bucket":"保单"},
        {"id":"as4","name":"保单抵押授权文件","severity":"blocking","bucket":"保单"},
        {"id":"as5","name":"担保人资料","severity":"blocking","bucket":"担保"},
        {"id":"as6","name":"董事/合伙人授权决议","severity":"blocking","bucket":"主体"},
    ],
    "SLL": [
        {"id":"s1","name":"SLL咨询申请表","severity":"blocking","bucket":"主体"},
        {"id":"s2","name":"公司注册证明（ACRA BizFile）","severity":"blocking","bucket":"主体"},
        {"id":"s3","name":"物业融资相关文件（参照BizProp checklist）","severity":"blocking","bucket":"物业"},
        {"id":"s4","name":"ESGpedia公司资料/碳排数据","severity":"blocking","bucket":"ESG"},
        {"id":"s5","name":"Sustainability Performance Targets（SPT）设定文件","severity":"blocking","bucket":"ESG"},
        {"id":"s6","name":"SPT达成证明材料（如适用）","severity":"advisory","bucket":"ESG"},
    ],
    "EFSGreen": [
        {"id":"e1","name":"EFS-Green申请表（常规版/Islamic版）","severity":"blocking","bucket":"主体"},
        {"id":"e2","name":"公司注册证明（ACRA BizFile）","severity":"blocking","bucket":"主体"},
        {"id":"e3","name":"本地股权证明（≥30%）","severity":"blocking","bucket":"主体"},
        {"id":"e4","name":"集团营业额证明（≤S$500m）","severity":"blocking","bucket":"财务"},
        {"id":"e5","name":"绿色项目/贸易用途说明书","severity":"blocking","bucket":"项目"},
        {"id":"e6","name":"项目合同/采购发票","severity":"blocking","bucket":"项目"},
        {"id":"e7","name":"现金流预测","severity":"blocking","bucket":"财务"},
        {"id":"e8","name":"绿色行业资质证明文件","severity":"blocking","bucket":"项目"},
    ],
}

BUILTIN_RULES = [
    {"id":"R01","name":"申请主体名称一致性","description":"所有文件中出现的公司名称必须与ACRA BizFile登记名称完全一致（允许'Pte Ltd'与'Private Limited'等价）","fields":["company_name","applicant_name"],"applies_to":"all","severity":"high"},
    {"id":"R02","name":"物业用途一致性","description":"估价报告、物业信息表、Title Deed三份文件中的物业用途代码必须一致，不可混用商业/工业/住宅","fields":["property_type","property_use"],"applies_to":["BizProp","BizPropPlus","SLL"],"severity":"high"},
    {"id":"R03","name":"MOF比率合规","description":"申请金额不得超过估价金额乘以产品MOF上限（BizProp 90%，BizPropPlus 130%，BizAssure 95%）","fields":["loan_amount","property_valuation"],"applies_to":["BizProp","BizPropPlus","BizAssure"],"severity":"high"},
    {"id":"R04","name":"本地股权比例达标","description":"股权文件中本地股权比例须满足产品要求（FlexiPay/BizGrow≥51%，BizAssist/EFSGreen≥30%）","fields":["local_shareholding"],"applies_to":["FlexiPay","BizGrow","BizAssist","EFSGreen"],"severity":"high"},
    {"id":"R05","name":"担保人与借款人不同主体","description":"担保人不可与贷款申请主体为同一法人实体","fields":["guarantor_name","borrower_name"],"applies_to":"all","severity":"high"},
    {"id":"R06","name":"银行流水账户与申请主体一致","description":"银行流水账户持有人须与申请公司名称一致","fields":["bank_account_holder","company_name"],"applies_to":["FlexiPay","BizGrow"],"severity":"medium"},
    {"id":"R07","name":"集团收入/雇员数达标（EFS WCL）","description":"BizAssist产品：财务报表须证明集团收入≤S$100m或雇员≤200人","fields":["group_revenue","employee_count"],"applies_to":["BizAssist"],"severity":"high"},
    {"id":"R08","name":"物业估价报告时效性","description":"估价报告出具日期须在申请日前12个月以内","fields":["valuation_date","application_date"],"applies_to":["BizProp","BizPropPlus","SLL"],"severity":"medium"},
    {"id":"R09","name":"保单被保险人与借款方关联","description":"人寿保单被保险人须为申请公司的Keyman或指定关键人","fields":["policy_insured","keyman_name"],"applies_to":["BizAssure"],"severity":"high"},
    {"id":"R10","name":"财务报表期间完整性","description":"提交的财务报表须覆盖申请日前至少2个完整会计年度","fields":["financial_period"],"applies_to":["BizGrow","BizAssist","BizProp","BizPropPlus"],"severity":"medium"},
]

# ─── Helpers ──────────────────────────────────────────────────────────────────

def extract_text_from_pdf(path: Path) -> tuple[str, float]:
    text_parts = []
    try:
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                t = page.extract_text()
                if t:
                    text_parts.append(t)
        text = "\n".join(text_parts)
        return (text, 95.0) if text.strip() else ("[Empty PDF]", 40.0)
    except Exception as e:
        return f"[PDF error: {e}]", 20.0

def extract_text_from_image(path: Path) -> tuple[str, float]:
    try:
        with open(path, "rb") as f:
            img_data = base64.standard_b64encode(f.read()).decode()
        ext = path.suffix.lower().lstrip(".")
        media_map = {"jpg":"image/jpeg","jpeg":"image/jpeg","png":"image/png","gif":"image/gif","webp":"image/webp"}
        media_type = media_map.get(ext, "image/jpeg")
        resp = client.chat.completions.create(
            model=MODEL, max_tokens=2000,
            messages=[{"role":"user","content":[
                {"type":"image_url","image_url":{"url":f"data:{media_type};base64,{img_data}"}},
                {"type":"text","text":"Extract ALL text from this image exactly as it appears. Preserve structure and numbers. Output only the extracted text."}
            ]}]
        )
        return resp.choices[0].message.content, 85.0
    except Exception as e:
        return f"[Image OCR error: {e}]", 0.0

def extract_file(path: Path) -> tuple[str, float]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(path)
    elif ext in [".jpg",".jpeg",".png",".gif",".webp",".tiff",".tif"]:
        return extract_text_from_image(path)
    elif ext in [".xlsx",".xls"]:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            rows = []
            for ws in wb.worksheets:
                rows.append(f"[Sheet: {ws.title}]")
                for row in ws.iter_rows(values_only=True):
                    rows.append("\t".join(str(c) if c is not None else "" for c in row))
            return "\n".join(rows), 99.0
        except Exception as e:
            return f"[Excel error: {e}]", 0.0
    elif ext == ".docx":
        try:
            from docx import Document as DocxDoc
            doc = DocxDoc(path)
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip()), 99.0
        except Exception as e:
            return f"[DOCX error: {e}]", 0.0
    elif ext in [".txt",".csv"]:
        try:
            return path.read_text(errors="replace"), 99.0
        except Exception as e:
            return f"[Read error: {e}]", 0.0
    return f"[Unsupported: {ext}]", 0.0

def load_custom_rules() -> list:
    f = RULES_DIR / "custom_rules.json"
    if f.exists():
        try:
            return json.loads(f.read_text())
        except:
            return []
    return []

# ─── Routes ───────────────────────────────────────────────────────────────────

@app.get("/api/products")
def get_products():
    return {"products":[
        {"key":"FlexiPay",    "label":"CIMB FlexiPay"},
        {"key":"BizGrow",     "label":"CIMB BizGrow / BizGrow-i"},
        {"key":"BizAssist",   "label":"CIMB BizAssist / BizAssist-i (EFS WCL)"},
        {"key":"BizProp",     "label":"CIMB BizProp / BizProp-i"},
        {"key":"BizPropPlus", "label":"CIMB BizProp Plus / BizProp Plus-i"},
        {"key":"BizAssure",   "label":"CIMB BizAssure"},
        {"key":"SLL",         "label":"SME Sustainability-Linked Loan"},
        {"key":"EFSGreen",    "label":"Enterprise Financing Scheme – Green / Green-i"},
    ]}

@app.get("/api/checklist/{product}")
def get_checklist(product: str):
    if product not in CHECKLISTS:
        raise HTTPException(404, f"Product '{product}' not found")
    return {"product": product, "items": CHECKLISTS[product]}

@app.get("/api/rules")
def get_rules():
    return {"builtin": BUILTIN_RULES, "custom": load_custom_rules()}

class SaveRulesRequest(BaseModel):
    rules: list

@app.post("/api/rules/custom")
def save_custom_rules(req: SaveRulesRequest):
    (RULES_DIR / "custom_rules.json").write_text(json.dumps(req.rules, ensure_ascii=False, indent=2))
    return {"saved": len(req.rules)}

@app.post("/api/upload")
async def upload_files(files: list[UploadFile] = File(...), product: str = Form(...)):
    if product not in CHECKLISTS:
        raise HTTPException(400, f"Unknown product: {product}")
    session_id  = str(uuid.uuid4())
    session_dir = UPLOAD_DIR / session_id
    session_dir.mkdir(parents=True)
    saved = []
    for f in files:
        dest    = session_dir / f.filename
        content = await f.read()
        dest.write_bytes(content)
        saved.append({"filename": f.filename, "size": len(content)})
    (session_dir / "meta.json").write_text(json.dumps({"product": product, "files": saved}, ensure_ascii=False))
    return {"session_id": session_id, "product": product, "files": saved}

@app.post("/api/ocr/{session_id}")
def run_ocr(session_id: str):
    session_dir = UPLOAD_DIR / session_id
    if not session_dir.exists():
        raise HTTPException(404, "Session not found")
    meta    = json.loads((session_dir / "meta.json").read_text())
    results = []
    for fi in meta["files"]:
        path = session_dir / fi["filename"]
        if not path.exists():
            continue
        text, conf = extract_file(path)
        results.append({"filename": fi["filename"], "confidence": round(conf,1),
                         "low_quality": conf < 80, "text": text, "char_count": len(text)})
        (session_dir / f"{fi['filename']}.txt").write_text(text, encoding="utf-8")
    (session_dir / "ocr_results.json").write_text(json.dumps({"results": results}, ensure_ascii=False))
    avg  = round(sum(r["confidence"] for r in results) / len(results), 1) if results else 0
    lows = sum(1 for r in results if r["low_quality"])
    return {"session_id": session_id, "total_files": len(results),
            "average_confidence": avg, "low_quality_count": lows, "results": results}

@app.post("/api/classify/{session_id}")
def run_classification(session_id: str):
    session_dir = UPLOAD_DIR / session_id
    if not session_dir.exists():
        raise HTTPException(404, "Session not found")
    meta      = json.loads((session_dir / "meta.json").read_text())
    ocr_path  = session_dir / "ocr_results.json"
    if not ocr_path.exists():
        raise HTTPException(400, "Run OCR first")
    ocr_data  = json.loads(ocr_path.read_text())
    product   = meta["product"]
    checklist = CHECKLISTS[product]
    summaries = [{"filename": r["filename"], "text_preview": r["text"][:800].replace("\n"," ")}
                 for r in ocr_data["results"]]
    prompt = f"""You are a CIMB SME loan document classification assistant.
Product: {product}
Checklist items:
{json.dumps(checklist, ensure_ascii=False, indent=2)}
Uploaded files with text previews:
{json.dumps(summaries, ensure_ascii=False, indent=2)}

For each uploaded file, determine which checklist item it best matches (if any).
Also extract key fields for contradiction detection.

Respond ONLY with valid JSON:
{{
  "matches": [
    {{
      "filename": "<filename>",
      "checklist_id": "<id or null>",
      "checklist_name": "<name or null>",
      "bucket": "<bucket or null>",
      "confidence": <0-100>,
      "reason": "<brief reason in Chinese>",
      "extracted_fields": {{
        "company_name": "<if found>",
        "property_type": "<if found>",
        "property_address": "<if found>",
        "loan_amount": "<if found>",
        "valuation_amount": "<if found>",
        "valuation_date": "<if found>",
        "local_shareholding": "<if found>",
        "guarantor_name": "<if found>",
        "bank_account_holder": "<if found>",
        "financial_period": "<if found>",
        "keyman_name": "<if found>",
        "application_date": "<if found>"
      }}
    }}
  ]
}}"""
    try:
        resp = client.chat.completions.create(
            model=MODEL, max_tokens=4000,
            response_format={"type":"json_object"},
            messages=[{"role":"user","content":prompt}]
        )
        raw = re.sub(r"^```(?:json)?\s*","",resp.choices[0].message.content.strip())
        raw = re.sub(r"\s*```$","",raw)
        classification = json.loads(raw)
    except Exception as e:
        raise HTTPException(500, f"Classification error: {e}")

    matches  = classification.get("matches",[])
    unmatched = [m["filename"] for m in matches if not m.get("checklist_id")]
    checklist_status = []
    for item in checklist:
        mf = next((m for m in matches if m.get("checklist_id")==item["id"]), None)
        status = ("matched" if mf and mf["confidence"]>=75 else
                  "low_confidence" if mf else "missing")
        checklist_status.append({**item, "status": status,
            "matched_file": mf["filename"] if mf else None,
            "confidence":   mf["confidence"] if mf else None})
    result = {
        "product": product, "matches": matches,
        "checklist_status": checklist_status,
        "unmatched_files": unmatched,
        "summary": {
            "matched":          sum(1 for s in checklist_status if s["status"]=="matched"),
            "low_confidence":   sum(1 for s in checklist_status if s["status"]=="low_confidence"),
            "missing_blocking": sum(1 for s in checklist_status if s["status"]=="missing" and s["severity"]=="blocking"),
            "missing_advisory": sum(1 for s in checklist_status if s["status"]=="missing" and s["severity"]=="advisory"),
            "unmatched_files":  len(unmatched),
        }
    }
    (session_dir / "classification.json").write_text(json.dumps(result, ensure_ascii=False))
    return result

@app.post("/api/detect/{session_id}")
def run_conflict_detection(session_id: str):
    session_dir = UPLOAD_DIR / session_id
    if not session_dir.exists():
        raise HTTPException(404, "Session not found")
    meta     = json.loads((session_dir / "meta.json").read_text())
    cls_path = session_dir / "classification.json"
    if not cls_path.exists():
        raise HTTPException(400, "Run classification first")
    cls_data = json.loads(cls_path.read_text())
    product  = meta["product"]
    all_fields: dict = {}
    for match in cls_data["matches"]:
        for k, v in match.get("extracted_fields", {}).items():
            if v:
                all_fields.setdefault(k, []).append({"value": v, "source": match["filename"]})
    custom_rules = load_custom_rules()
    all_rules    = BUILTIN_RULES + custom_rules
    applicable   = [r for r in all_rules
                    if r.get("applies_to")=="all" or product in (r.get("applies_to") or [])]
    doc_context  = [{"filename": m["filename"], "checklist_name": m.get("checklist_name"),
                     "extracted_fields": m.get("extracted_fields",{})} for m in cls_data["matches"]]
    prompt = f"""You are a CIMB SME loan document contradiction detection assistant.
Product: {product}
Applicable rules:
{json.dumps(applicable, ensure_ascii=False, indent=2)}
Extracted fields across all documents:
{json.dumps(all_fields, ensure_ascii=False, indent=2)}
Document context:
{json.dumps(doc_context, ensure_ascii=False, indent=2)}

Identify contradictions based on the rules. For each finding specify rule, files, conflicting values, explanation.

Respond ONLY with valid JSON:
{{
  "findings": [
    {{
      "rule_id": "<id>",
      "rule_name": "<name>",
      "severity": "high|medium|low",
      "field": "<field>",
      "files_involved": ["<file1>","<file2>"],
      "value_a": {{"source":"<file>","value":"<val>"}},
      "value_b": {{"source":"<file>","value":"<val>"}},
      "explanation": "<Chinese explanation>",
      "recommendation": "<Chinese recommendation>"
    }}
  ],
  "rules_checked": <number>,
  "passed": <number>,
  "conclusion": "<Chinese conclusion 2-3 sentences>"
}}"""
    try:
        resp = client.chat.completions.create(
            model=MODEL, max_tokens=4000,
            response_format={"type":"json_object"},
            messages=[{"role":"user","content":prompt}]
        )
        raw = re.sub(r"^```(?:json)?\s*","",resp.choices[0].message.content.strip())
        raw = re.sub(r"\s*```$","",raw)
        detection = json.loads(raw)
    except Exception as e:
        raise HTTPException(500, f"Detection error: {e}")
    (session_dir / "detection.json").write_text(json.dumps(detection, ensure_ascii=False))
    return detection

@app.get("/api/session/{session_id}/text/{filename}")
def get_file_text(session_id: str, filename: str):
    path = UPLOAD_DIR / session_id / f"{filename}.txt"
    if not path.exists():
        raise HTTPException(404, "Text not found")
    return {"text": path.read_text(encoding="utf-8")}

@app.post("/api/session/{session_id}/text/{filename}")
async def update_file_text(session_id: str, filename: str, request: Request):
    body = await request.json()
    path = UPLOAD_DIR / session_id / f"{filename}.txt"
    path.write_text(body.get("text",""), encoding="utf-8")
    ocr_path = UPLOAD_DIR / session_id / "ocr_results.json"
    if ocr_path.exists():
        data = json.loads(ocr_path.read_text())
        for r in data["results"]:
            if r["filename"] == filename:
                r["text"] = body.get("text","")
                r["manually_corrected"] = True
        ocr_path.write_text(json.dumps(data, ensure_ascii=False))
    return {"saved": True}

@app.post("/api/session/{session_id}/assign")
async def assign_unmatched(session_id: str, request: Request):
    body     = await request.json()
    cls_path = UPLOAD_DIR / session_id / "classification.json"
    if not cls_path.exists():
        raise HTTPException(400, "No classification data")
    cls_data    = json.loads(cls_path.read_text())
    filename    = body["filename"]
    checklist_id = body.get("checklist_id")
    meta        = json.loads((UPLOAD_DIR / session_id / "meta.json").read_text())
    for match in cls_data["matches"]:
        if match["filename"] == filename:
            match["checklist_id"] = checklist_id
            match["manually_assigned"] = True
            if checklist_id:
                item = next((c for c in CHECKLISTS[meta["product"]] if c["id"]==checklist_id), None)
                if item:
                    match["checklist_name"] = item["name"]
                    match["bucket"]         = item["bucket"]
                    match["confidence"]     = 100
    cls_data["unmatched_files"] = [f for f in cls_data["unmatched_files"] if f != filename]
    cls_path.write_text(json.dumps(cls_data, ensure_ascii=False))
    return {"saved": True}

@app.get("/health")
def health():
    return {"status": "ok"}

# ─── Serve frontend (must be last) ───────────────────────────────────────────
if FRONTEND.exists():
    app.mount("/", StaticFiles(directory=str(FRONTEND), html=True), name="static")
